from io import BytesIO
from pathlib import Path
from importlib.util import find_spec
from unittest.mock import patch

import numpy as np
import openpyxl
from streamlit.testing.v1 import AppTest

from portfolio.bundle import read_bundle
from portfolio.calibration import fit_ff5
from portfolio.export import workbook_bytes
from portfolio.results import research_run
from test_calibration import sample_bundle


ROOT = Path(__file__).resolve().parents[1]


def test_missing_sklearn_explains_deployment_failure_before_calculating():
    with patch("portfolio.deployment.find_spec", side_effect=lambda name: None if name == "sklearn" else find_spec(name)):
        app = AppTest.from_file(ROOT / "ui/legacy_app.py", default_timeout=30).run()
    assert not app.exception
    assert "scikit-learn" in app.error[0].value
    assert any("Reboot app" in item.value for item in app.markdown)
    assert not app.metric
    assert "active_run" not in app.session_state


def test_slider_uses_solved_weights_without_rebuilding_and_failure_keeps_last_run():
    app = AppTest.from_file(ROOT / "ui/legacy_app.py", default_timeout=30).run()
    assert not app.exception
    assert "Legacy demonstration" in app.warning[0].value
    first = app.session_state["active_run"]
    with patch("portfolio.results.efficient_frontier", side_effect=AssertionError("Slider recalculated the frontier")):
        app.slider[0].set_value(100).run()
    assert not app.exception
    assert app.session_state["active_run"].run_id == first.run_id
    assert app.metric[0].value == f"{first.frontier.returns[-1] * 12:.2%}"
    app.sidebar.number_input[0].set_value(1.0)
    app.sidebar.button[0].click().run()
    assert not app.exception
    assert "infeasible" in app.error[0].value
    assert app.session_state["active_run"].run_id == first.run_id


def test_single_fund_and_missing_research_bundle_are_handled():
    app = AppTest.from_file(ROOT / "ui/legacy_app.py", default_timeout=30).run()
    app.sidebar.multiselect[0].set_value(["ACWI"])
    app.sidebar.button[0].click().run()
    assert not app.exception
    assert not app.slider
    assert len(app.session_state["active_run"].frontier.weights) == 1
    app.sidebar.radio[0].set_value("Return-based FF5").run()
    assert not app.exception
    assert any("research data is still needed" in x.value for x in app.subheader)
    assert not app.metric


def test_ff5_workflow_and_export_preserve_selected_portfolio_and_assumptions():
    raw, _ = sample_bundle()
    data = read_bundle(raw)
    fit = fit_ff5(data.returns.tail(120), data.factors, data.cash)
    run = research_run(data, fit, ["A", "B"], 1.0, [0.06, 0.02, 0.03, 0.015, 0.01], 0.03, 0.5)
    result = workbook_bytes(run, 50, investment=10000)
    workbook = openpyxl.load_workbook(BytesIO(result), data_only=False)
    assert "Monthly covariance" in workbook.sheetnames
    assert "Wealth illustration" in workbook.sheetnames
    rows = list(workbook["Portfolio"].values)
    weight_column = rows[0].index("Weight")
    np.testing.assert_allclose([row[weight_column] for row in rows[1:]], run.frontier.weights[50])
    assert all(cell.data_type != "f" for sheet in workbook for row in sheet for cell in row)
    assert run.settings["premiums_annual_arithmetic"] == [0.06, 0.01, 0.015, 0.0075, 0.005]


def test_research_interface_accepts_a_bundle_and_runs_premium_sensitivity():
    raw, _ = sample_bundle()
    app = AppTest.from_file(ROOT / "ui/legacy_app.py", default_timeout=30).run()
    with patch("streamlit.delta_generator.DeltaGenerator.file_uploader", return_value=BytesIO(raw)):
        app.sidebar.radio[0].set_value("Return-based FF5").run()
        assert not app.exception
        assert app.session_state["active_run"].mode == "Return-based FF5"
        app.sidebar.selectbox[1].set_value(0.0)
        app.sidebar.button[0].click().run()
        assert not app.exception
        assert app.session_state["active_run"].settings["premiums_annual_arithmetic"][1:] == [0.0] * 4
        app.checkbox[0].check().run()
        assert not app.exception
